# -*- coding: utf-8 -*-
"""Đọc file đăng ký tải về từ web và đối chiếu với các file "Bao cao 1".

Web chỉ thu thập nguyện vọng: học viên tự gõ họ tên, ngày sinh, CCCD, khóa cũ.
Dữ liệu đưa vào văn bản chính thức thì phải lấy từ "Bao cao 1" — địa chỉ, số
GPLX, mã học viên, mã khóa gốc đều nằm ở đó. Vì vậy bước này chỉ dùng bản đăng
ký để CHỌN người, còn nội dung thì lấy nguyên từ hồ sơ khóa cũ.
"""

import csv
import re
from dataclasses import dataclass
from pathlib import Path

from .vn import bo_dau

# tên cột chấp nhận được, đã bỏ dấu và viết hoa
_TEN_COT = {
    "hoTen": ("HO VA TEN", "HOTEN", "HO TEN"),
    "ngaySinh": ("NGAY SINH", "NGAYSINH"),
    "cccd": ("CCCD", "SO CCCD", "CMND", "CAN CUOC"),
    "diaChi": ("NOI THUONG TRU", "DIA CHI", "DIACHI"),
    "khoaGoc": ("KHOA DA HOC", "KHOAGOC", "KHOA CU", "KHOA GOC"),
    "khoaDich": ("DOT GHEP VAO", "KHOADICH", "KHOA DICH", "DOT GHEP"),
    "soDienThoai": ("DIEN THOAI", "SODIENTHOAI", "SO DIEN THOAI", "SDT"),
    "trangThai": ("TRANG THAI", "TRANGTHAI"),
    "ghiChu": ("GHI CHU", "GHICHU"),
}


@dataclass
class NguoiDangKy:
    ho_ten: str
    ngay_sinh: str
    cccd: str
    khoa_goc: str
    khoa_dich: str
    so_dien_thoai: str = ""
    dia_chi: str = ""
    trang_thai: str = ""
    ghi_chu: str = ""
    dong: int = 0


def _chuan_cot(s):
    return re.sub(r"\s+", " ", bo_dau(str(s or "")).strip())


def _do_cot(hang_tieu_de):
    """Ánh xạ tên cột trong file về khóa nội bộ."""
    anh_xa = {}
    for i, o in enumerate(hang_tieu_de):
        c = _chuan_cot(o)
        for khoa, ten in _TEN_COT.items():
            if c in ten:
                anh_xa[khoa] = i
                break
    return anh_xa


def _cac_hang(duong_dan):
    p = Path(duong_dan)
    duoi = p.suffix.lower()
    if duoi == ".csv":
        # utf-8-sig: file web xuất ra có BOM để Excel đọc đúng tiếng Việt
        with p.open(encoding="utf-8-sig", newline="") as f:
            return [list(r) for r in csv.reader(f)]
    if duoi == ".xls":
        import xlrd
        sh = xlrd.open_workbook(str(p)).sheet_by_index(0)
        return [[sh.cell(r, c).value for c in range(sh.ncols)] for r in range(sh.nrows)]
    import openpyxl
    ws = openpyxl.load_workbook(str(p), data_only=True).worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _o(v):
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        v = int(v)
    return str(v).strip()


def doc(duong_dan, chi_duyet=True):
    """Đọc file đăng ký. Trả về (danh sách NguoiDangKy, danh sách cảnh báo)."""
    hang = _cac_hang(duong_dan)
    if not hang:
        raise SystemExit(f"{Path(duong_dan).name}: file rỗng")

    anh_xa, dong_tieu_de = {}, -1
    for i, h in enumerate(hang[:10]):
        thu = _do_cot(h)
        if {"hoTen", "cccd"} <= set(thu):
            anh_xa, dong_tieu_de = thu, i
            break
    if dong_tieu_de < 0:
        raise SystemExit(
            f"{Path(duong_dan).name}: không tìm thấy dòng tiêu đề có cột 'Họ và Tên' và 'CCCD'"
        )

    ds, canh_bao = [], []
    for i, h in enumerate(hang[dong_tieu_de + 1:], start=dong_tieu_de + 2):
        lay = lambda k: _o(h[anh_xa[k]]) if k in anh_xa and anh_xa[k] < len(h) else ""  # noqa: E731
        ten = lay("hoTen")
        if not ten:
            continue
        n = NguoiDangKy(
            ho_ten=re.sub(r"\s+", " ", ten).upper(),
            ngay_sinh=lay("ngaySinh"),
            cccd=re.sub(r"\D", "", lay("cccd")),
            khoa_goc=re.sub(r"\s+", "", lay("khoaGoc")).upper(),
            khoa_dich=re.sub(r"\s+", "", lay("khoaDich")).upper(),
            so_dien_thoai=lay("soDienThoai"),
            dia_chi=lay("diaChi"),
            trang_thai=lay("trangThai"),
            ghi_chu=lay("ghiChu"),
            dong=i,
        )
        if chi_duyet and n.trang_thai and bo_dau(n.trang_thai) not in ("DA DUYET", "DUYET"):
            canh_bao.append(
                f"BỎ QUA dòng {i} ({n.ho_ten}): trạng thái '{n.trang_thai}' chưa được duyệt."
            )
            continue
        ds.append(n)
    return ds, canh_bao


def doi_chieu(nguoi_dang_ky, bao_caos):
    """Khớp từng người đăng ký với học viên trong các file "Bao cao 1".

    Ưu tiên CCCD (định danh thật); nếu không có mới dò theo tên + ngày sinh.
    Trả về (danh sách HocVien lấy từ Bao cao 1, danh sách cảnh báo).
    """
    theo_cccd, theo_ten_ns, theo_ten = {}, {}, {}
    dat = set()
    for bc in bao_caos:
        for hv in bc.hoc_vien:
            if hv.cccd:
                theo_cccd.setdefault(hv.cccd, []).append(hv)
            theo_ten_ns.setdefault((bo_dau(hv.ho_ten), hv.ngay_sinh), []).append(hv)
            theo_ten.setdefault(bo_dau(hv.ho_ten), []).append(hv)
            if hv.cccd in bc.cccd_dat:
                dat.add(hv.cccd)

    ket_qua, canh_bao, da_lay = [], [], set()
    for n in nguoi_dang_ky:
        ung = theo_cccd.get(n.cccd, [])
        cach = "CCCD"
        if not ung:
            ung = theo_ten_ns.get((bo_dau(n.ho_ten), n.ngay_sinh), [])
            cach = "tên + ngày sinh"
        if not ung:
            ung = theo_ten.get(bo_dau(n.ho_ten), [])
            cach = "chỉ tên"

        if not ung:
            canh_bao.append(
                f"KHÔNG TÌM THẤY: {n.ho_ten} — {n.cccd} (đăng ký dòng {n.dong}, khai khóa "
                f"{n.khoa_goc or '?'}). Không có trong file Bao cao 1 nào."
            )
            continue

        if len(ung) > 1:
            # nhiều người trùng: thử thu hẹp bằng khóa học viên tự khai
            hep = [h for h in ung if h.ma_khoa_goc == n.khoa_goc] if n.khoa_goc else []
            if len(hep) == 1:
                ung = hep
            else:
                canh_bao.append(
                    f"TRÙNG NHIỀU HỒ SƠ: {n.ho_ten} — {n.cccd} khớp {len(ung)} dòng "
                    + " | ".join(f"{h.ma_khoa_goc}/{h.nguon} dòng {h.dong}" for h in ung)
                    + " — hãy sửa mã khóa trong bản đăng ký rồi chạy lại."
                )
                continue

        hv = ung[0]

        if cach != "CCCD":
            canh_bao.append(
                f"KHỚP MỜ ({cach}): {n.ho_ten} — CCCD đăng ký '{n.cccd}' không có trong hồ sơ, "
                f"đã dùng {hv.ma_khoa_goc}/{hv.nguon} dòng {hv.dong} (CCCD hồ sơ '{hv.cccd}'). "
                "Kiểm tra lại trước khi in."
            )
        if hv.cccd in dat:
            canh_bao.append(
                f"ĐÃ ĐẠT RỒI: {hv.ho_ten} ({hv.ma_khoa_goc}) nằm trong khối ĐẠT của sheet bc2 "
                "— vẫn đăng ký thi ghép. Kiểm tra lại."
            )
        if n.khoa_goc and n.khoa_goc != hv.ma_khoa_goc:
            canh_bao.append(
                f"LỆCH KHÓA: {hv.ho_ten} tự khai khóa '{n.khoa_goc}' nhưng hồ sơ ở "
                f"'{hv.ma_khoa_goc}' — lấy theo hồ sơ."
            )
        if hv.cccd in da_lay:
            canh_bao.append(f"ĐĂNG KÝ TRÙNG: {hv.ho_ten} — {hv.cccd} xuất hiện nhiều lần, chỉ lấy 1.")
            continue

        da_lay.add(hv.cccd)
        ket_qua.append(hv)

    return ket_qua, canh_bao
